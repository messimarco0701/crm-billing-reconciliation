"""
Reconciles CRM deals against billing contracts and billing plans.

The flow is deliberately simple:
  1. SQL staging normalizes names and VAT numbers (sql/01_staging.sql)
  2. SQL matches deals to contracts on VAT (sql/02_match_on_vat.sql)
  3. Python matches what is left on company name, with a fuzzy score
  4. SQL runs one check per business rule (sql/03_checks.sql)
  5. Results go to an Excel file that a non technical colleague can filter

Run it with:  python src/reconcile.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import duckdb
import pandas as pd
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz, process

ROOT = Path(__file__).resolve().parents[1]
SQL_DIR = ROOT / "sql"
SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}


def load_rules(path: Path = ROOT / "rules.yaml") -> dict:
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def run_sql(con: duckdb.DuckDBPyConnection, name: str, **params) -> None:
    sql = (SQL_DIR / name).read_text(encoding="utf-8")
    for key, value in params.items():
        sql = sql.replace("{" + key + "}", str(value))
    con.execute(sql)


def fuzzy_match(con: duckdb.DuckDBPyConnection, threshold: float) -> pd.DataFrame:
    """Match deals without a usable VAT to contracts nobody has claimed yet."""
    deals = con.execute("""
        SELECT deal_id, name_key FROM stg_deals
        WHERE vat IS NULL AND deal_id NOT IN (SELECT deal_id FROM matches)
    """).df()
    candidates = con.execute("""
        SELECT contract_id, name_key FROM stg_contracts
        WHERE contract_id NOT IN (SELECT contract_id FROM matches)
    """).df()

    found = []
    pool = dict(zip(candidates.contract_id, candidates.name_key))
    for deal_id, key in zip(deals.deal_id, deals.name_key):
        if not pool:
            break
        best = process.extract(key, pool, scorer=fuzz.token_sort_ratio, limit=2)
        top_name, top_score, top_id = best[0]
        runner_up = best[1][1] if len(best) > 1 else 0
        # Accept only a clear winner: two similar candidates means a human should decide
        if top_score >= threshold and top_score - runner_up >= 5:
            found.append({"deal_id": deal_id, "contract_id": top_id,
                          "match_method": "fuzzy_name", "match_score": round(top_score, 1)})
            pool.pop(top_id)

    result = pd.DataFrame(found, columns=["deal_id", "contract_id", "match_method", "match_score"])
    if not result.empty:
        con.register("fuzzy_result", result)
        con.execute("INSERT INTO matches SELECT * FROM fuzzy_result")
    return result


def run(data_dir: Path = ROOT / "data", rules: dict | None = None) -> dict[str, pd.DataFrame]:
    rules = rules or load_rules()
    con = duckdb.connect()
    con.execute("SET VARIABLE won_stage = ?", [rules["won_stage"]])
    con.execute("SET VARIABLE amount_tolerance = ?", [rules["amount_tolerance_pct"]])
    con.execute("SET VARIABLE minimum_fee = ?", [rules["minimum_annual_fee"]])
    con.execute("SET VARIABLE deal_name_pattern = ?", [rules["deal_name_pattern"]])

    run_sql(con, "01_staging.sql", data_dir=data_dir.as_posix())
    run_sql(con, "02_match_on_vat.sql")
    fuzzy_match(con, rules["fuzzy_match_threshold"])
    run_sql(con, "03_checks.sql")

    anomalies = con.execute("SELECT * FROM anomalies").df()
    anomalies["_sev"] = anomalies["severity"].map(SEVERITY_ORDER)
    anomalies = (anomalies.sort_values(["_sev", "amount_at_stake"], ascending=[True, False])
                 .drop(columns="_sev").reset_index(drop=True))

    summary = (anomalies.groupby(["anomaly_type", "severity"], as_index=False)
               .agg(cases=("anomaly_type", "size"), amount_at_stake=("amount_at_stake", "sum")))
    summary["_sev"] = summary["severity"].map(SEVERITY_ORDER)
    summary = summary.sort_values(["_sev", "cases"], ascending=[True, False]).drop(columns="_sev")

    coverage = con.execute("""
        SELECT
            (SELECT count(*) FROM stg_deals)                                    AS deals,
            (SELECT count(*) FROM stg_contracts)                                AS contracts,
            (SELECT count(*) FROM matches WHERE match_method = 'vat')           AS matched_on_vat,
            (SELECT count(*) FROM matches WHERE match_method = 'fuzzy_name')    AS matched_on_name
    """).df()

    matches = con.execute("""
        SELECT m.*, d.company_name AS crm_name, c.customer_name AS billing_name
        FROM matches m
        JOIN stg_deals d USING (deal_id)
        JOIN stg_contracts c USING (contract_id)
        ORDER BY m.match_method DESC, m.match_score
    """).df()

    return {"summary": summary, "anomalies": anomalies, "matches": matches, "coverage": coverage}


def write_excel(results: dict[str, pd.DataFrame], rules: dict, path: Path) -> None:
    path.parent.mkdir(exist_ok=True)
    rules_df = pd.DataFrame({"rule": list(rules), "value": [str(v) for v in rules.values()]})
    sheets = {
        "Summary": results["summary"],
        "Anomalies": results["anomalies"],
        "Matches": results["matches"],
        "Rules": rules_df,
    }
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for i, col in enumerate(df.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 70)


def main() -> None:
    rules = load_rules()
    if not (ROOT / "data" / "crm_deals.csv").exists():
        sys.exit("No data found. Run  python src/generate_data.py  first.")

    results = run(rules=rules)
    out = ROOT / "reports" / "reconciliation_report.xlsx"
    write_excel(results, rules, out)

    cov = results["coverage"].iloc[0]
    print(f"Deals: {cov.deals}   Contracts: {cov.contracts}   "
          f"Matched on VAT: {cov.matched_on_vat}   Matched on name: {cov.matched_on_name}\n")
    print(results["summary"].to_string(index=False, float_format=lambda x: f"{x:,.2f}"))
    print(f"\nTotal cases to review: {len(results['anomalies'])}")
    print(f"Report written to {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
